"""
Comptes chauffeurs en masse : import Excel, envoi des identifiants par SMS.

Chaque compte importé reçoit le mot de passe initial commun
(CHAUFFEUR_MOT_DE_PASSE_INITIAL, « Camusat@2026 » par défaut) et doit le changer
à sa première connexion.
"""
import io
import unicodedata
from datetime import datetime, timezone

import httpx
import pandas as pd
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session

from ..config import settings
from ..database import get_db
from ..models.user import User
from ..models.vehicule import Vehicule
from ..services.auth_service import hash_password, require_admin
from ..services.sms import ErreurSMS, client_sms, envoyer_sms, message_identifiants, normaliser_telephone, sms_actif

router = APIRouter(prefix="/api/auth/chauffeurs", tags=["Comptes chauffeurs"])

COLONNES = ["Nom complet", "Téléphone", "Filiale", "Immatriculation", "Identifiant", "Email"]
MAX_LIGNES = 2000


class LigneImport(BaseModel):
    ligne: int
    nom: str | None
    username: str | None
    telephone: str | None
    statut: str               # CREE | IGNORE | ERREUR
    message: str | None = None
    sms: str | None = None    # ENVOYE | ECHEC | NON_CONFIGURE | SANS_NUMERO | None (non demandé)
    user_id: int | None = None


class ResultatImport(BaseModel):
    crees: int
    ignores: int
    erreurs: int
    sms_envoyes: int
    sms_echecs: int
    sms_actif: bool
    mot_de_passe_initial: str
    app_url: str
    lignes: list[LigneImport]


class ResultatIdentifiants(BaseModel):
    username: str
    telephone: str | None
    mot_de_passe_initial: str
    sms: str
    message: str | None = None


def _cle(texte: str) -> str:
    """En-tête normalisé : minuscules, sans accents ni espaces superflus."""
    t = unicodedata.normalize("NFD", str(texte)).encode("ascii", "ignore").decode()
    return " ".join(t.lower().replace("_", " ").replace("-", " ").split())


ALIAS = {
    "nom": {"nom complet", "nom et prenom", "prenom et nom", "nom prenom", "prenom nom", "nom", "chauffeur", "full name"},
    "telephone": {"telephone", "tel", "numero", "numero de telephone", "portable", "mobile", "contact", "phone"},
    "filiale": {"filiale", "entite", "societe"},
    "plaque": {"immatriculation", "plaque", "plaque immatriculation", "matricule", "vehicule"},
    "username": {"identifiant", "username", "login", "nom d utilisateur"},
    "email": {"email", "e mail", "mail", "adresse email"},
}


def _texte(val) -> str | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    t = str(val).strip()
    if t.endswith(".0") and t[:-2].isdigit():
        t = t[:-2]
    return t or None


def _envoyer(user: User, client: httpx.Client) -> tuple[str, str | None]:
    if not user.telephone:
        return "SANS_NUMERO", "Aucun numéro de téléphone"
    if not sms_actif():
        return "NON_CONFIGURE", "SMS non configurés sur le serveur"
    try:
        envoyer_sms(user.telephone, message_identifiants(user.username, settings.chauffeur_mot_de_passe_initial), client)
        user.identifiants_envoyes_at = datetime.now(timezone.utc)
        return "ENVOYE", None
    except ErreurSMS as e:
        return "ECHEC", str(e)


@router.get("/modele")
def modele_excel(_: User = Depends(require_admin)):
    """Fichier Excel à remplir pour l'import en masse."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Chauffeurs"
    ws.append(COLONNES)
    ws.append(["Moussa Diop", "77 123 45 67", "Camusat Sénégal", "AA 028 TQ", "", ""])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003C71")
    for col, largeur in zip("ABCDEF", (28, 18, 22, 18, 18, 28)):
        ws.column_dimensions[col].width = largeur
    aide = wb.create_sheet("Aide")
    for ligne in [
        ["Colonne", "Obligatoire", "Explication"],
        ["Nom complet", "Oui", "Prénom et nom du chauffeur"],
        ["Téléphone", "Oui", "Reçoit les identifiants par SMS. Ex : 77 123 45 67 ou +221771234567"],
        ["Filiale", "Oui*", "*Sinon la filiale choisie dans la fenêtre d'import est utilisée"],
        ["Immatriculation", "Non", "Véhicule attribué ; doit exister dans la flotte"],
        ["Identifiant", "Non", "Par défaut : le numéro de téléphone (9 chiffres, ex : 771234567)"],
        ["Email", "Non", ""],
    ]:
        aide.append(ligne)
    for cell in aide[1]:
        cell.font = Font(bold=True)
    aide.column_dimensions["A"].width = 18
    aide.column_dimensions["C"].width = 70
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modele_import_chauffeurs.xlsx"'},
    )


@router.post("/import", response_model=ResultatImport)
async def importer(
    file: UploadFile = File(...),
    filiale_defaut: str = Form(""),
    envoyer: bool = Form(False),
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    contenu = await file.read()
    try:
        if (file.filename or "").lower().endswith(".csv"):
            df = pd.read_csv(io.BytesIO(contenu), dtype=str, sep=None, engine="python")
        else:
            df = pd.read_excel(io.BytesIO(contenu), dtype=str)
    except Exception:
        raise HTTPException(400, "Fichier illisible : utilisez le modèle Excel (.xlsx)")

    colonnes: dict[str, str] = {}
    for col in df.columns:
        cle = _cle(col)
        for champ, alias in ALIAS.items():
            if cle in alias and champ not in colonnes:
                colonnes[champ] = col
    manquantes = [n for n, c in (("Nom complet", "nom"), ("Téléphone", "telephone")) if c not in colonnes]
    if manquantes:
        raise HTTPException(400, f"Colonnes manquantes : {', '.join(manquantes)}. Utilisez le modèle Excel.")
    if len(df) > MAX_LIGNES:
        raise HTTPException(400, f"Fichier trop long ({len(df)} lignes, maximum {MAX_LIGNES})")

    plaques = {p for (p,) in db.query(Vehicule.plaque_immatriculation).all()}
    plaques_norm = {p.replace(" ", "").replace("-", "").upper(): p for p in plaques}
    noms_pris = {u.lower() for (u,) in db.query(User.username).all()}
    tels_pris = {t for (t,) in db.query(User.telephone).filter(User.telephone.isnot(None)).all()}
    emails_pris = {e.lower() for (e,) in db.query(User.email).filter(User.email.isnot(None)).all()}
    mdp_hash = hash_password(settings.chauffeur_mot_de_passe_initial)
    filiale_defaut = filiale_defaut.strip()

    lignes: list[LigneImport] = []
    a_creer: list[tuple[LigneImport, User]] = []
    for i, row in df.iterrows():
        n = int(i) + 2   # numéro de ligne dans Excel (en-tête = ligne 1)
        val = lambda champ: _texte(row[colonnes[champ]]) if champ in colonnes else None
        nom, tel_brut = val("nom"), val("telephone")
        if not nom and not tel_brut and not val("username"):
            continue     # ligne vide
        res = LigneImport(ligne=n, nom=nom, username=None, telephone=tel_brut, statut="ERREUR")
        lignes.append(res)

        if not nom:
            res.message = "Nom complet manquant"; continue
        tel = normaliser_telephone(tel_brut)
        if not tel:
            res.message = f"Téléphone invalide : {tel_brut or 'vide'}"; continue
        res.telephone = tel
        filiale = val("filiale") or filiale_defaut
        if not filiale:
            res.message = "Filiale manquante (colonne ou filiale par défaut)"; continue
        username = (val("username") or tel.removeprefix("+221")).strip()
        res.username = username
        if username.lower() in noms_pris:
            res.statut, res.message = "IGNORE", "Identifiant déjà utilisé (compte existant)"; continue
        if tel in tels_pris:
            res.statut, res.message = "IGNORE", "Téléphone déjà associé à un compte"; continue
        email = val("email")
        if email and email.lower() in emails_pris:
            res.statut, res.message = "IGNORE", "Email déjà utilisé"; continue

        plaque, avertissement = None, None
        if val("plaque"):
            plaque = plaques_norm.get(val("plaque").replace(" ", "").replace("-", "").upper())
            if not plaque:
                avertissement = f"Véhicule {val('plaque')} absent de la flotte : non attribué"

        user = User(
            username=username, full_name=nom, email=email, telephone=tel,
            hashed_password=mdp_hash, role="CHAUFFEUR", is_active=True,
            vehicule_plaque=plaque, filiale=filiale, doit_changer_mdp=True,
        )
        noms_pris.add(username.lower()); tels_pris.add(tel)
        if email:
            emails_pris.add(email.lower())
        res.statut, res.message = "CREE", avertissement
        a_creer.append((res, user))

    for _res, user in a_creer:
        db.add(user)
    db.flush()
    for res, user in a_creer:
        res.user_id = user.id
    db.commit()

    if envoyer and a_creer:
        with client_sms() as client:
            for res, user in a_creer:
                res.sms, erreur = _envoyer(user, client)
                if erreur and res.sms == "ECHEC":
                    res.message = f"{res.message + ' · ' if res.message else ''}SMS : {erreur}"
        db.commit()

    return ResultatImport(
        crees=sum(l.statut == "CREE" for l in lignes),
        ignores=sum(l.statut == "IGNORE" for l in lignes),
        erreurs=sum(l.statut == "ERREUR" for l in lignes),
        sms_envoyes=sum(l.sms == "ENVOYE" for l in lignes),
        sms_echecs=sum(l.sms == "ECHEC" for l in lignes),
        sms_actif=sms_actif(),
        mot_de_passe_initial=settings.chauffeur_mot_de_passe_initial,
        app_url=settings.app_url,
        lignes=lignes,
    )


@router.post("/{user_id}/identifiants", response_model=ResultatIdentifiants)
def renvoyer_identifiants(user_id: int, db: Session = Depends(get_db), _: User = Depends(require_admin)):
    """Remet le mot de passe initial (à changer à la connexion) et l'envoie par SMS.
    Sert aussi en cas de mot de passe oublié."""
    user = db.query(User).filter(User.id == user_id, User.role == "CHAUFFEUR").first()
    if not user:
        raise HTTPException(404, "Chauffeur introuvable")
    user.hashed_password = hash_password(settings.chauffeur_mot_de_passe_initial)
    user.doit_changer_mdp = True
    with client_sms() as client:
        statut, erreur = _envoyer(user, client)
    db.commit()
    return ResultatIdentifiants(
        username=user.username, telephone=user.telephone,
        mot_de_passe_initial=settings.chauffeur_mot_de_passe_initial,
        sms=statut, message=erreur,
    )


@router.get("/sms")
def statut_sms(_: User = Depends(require_admin)):
    return {
        "actif": sms_actif(), "app_url": settings.app_url,
        "mot_de_passe_initial": settings.chauffeur_mot_de_passe_initial,
        "signature": settings.orange_sms_signature,
    }


class TestSMS(BaseModel):
    telephone: str


@router.post("/sms/test")
def tester_sms(data: TestSMS, _: User = Depends(require_admin)):
    """Envoie un SMS de test pour vérifier la configuration Orange SMS Pro."""
    tel = normaliser_telephone(data.telephone)
    if not tel:
        raise HTTPException(400, "Numéro de téléphone invalide")
    try:
        reponse = envoyer_sms(tel, "eFlotte CAMUSAT: SMS de test. La configuration des SMS fonctionne.")
    except ErreurSMS as e:
        raise HTTPException(502, str(e))
    return {"telephone": tel, "reponse": reponse}
